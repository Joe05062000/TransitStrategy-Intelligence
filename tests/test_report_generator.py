"""Tests for the PowerPoint Executive Report Generator (src/report_generator.py).

Verifies the generated file is a structurally valid .pptx (opens with
python-pptx), has the expected number of slides, and surfaces both the
REAL DATA and HYPOTHETICAL labels -- without any AI/network dependency.
"""

import io
import sys
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from diagnostics import run_diagnostics  # noqa: E402
from kpi import TARGET_AGENCIES, agency_year_kpis, mode_group_year_kpis  # noqa: E402
from report_generator import build_report  # noqa: E402
from scenarios import project_scenario, summarize_scenario, suggest_reference_assumptions  # noqa: E402


@pytest.fixture(scope="module")
def built_report():
    agency = "LA Metro"
    agency_kpis = agency_year_kpis()
    mode_kpis = mode_group_year_kpis()
    hist = agency_kpis[agency_kpis["agency"] == agency].sort_values("year")

    assumptions = suggest_reference_assumptions(hist, lookback_years=5)
    projection = project_scenario(hist, assumptions)
    summary = summarize_scenario(projection)

    findings = run_diagnostics(agency_kpis, list(TARGET_AGENCIES.values()), 2015, 2025)
    bench = mode_kpis[(mode_kpis["mode_group"] == "Rail") & (mode_kpis["agency"].isin(TARGET_AGENCIES.values()))]

    buf = build_report(
        agency=agency,
        report_year=2025,
        agency_history=hist,
        mode_bench=bench,
        mode_group_label="Rail",
        peer_agencies=["Sacramento RT", "BART"],
        findings=findings,
        assumptions=assumptions,
        projection=projection,
        scenario_summary=summary,
    )
    return buf


def test_report_is_nonempty_bytes(built_report):
    data = built_report.getvalue()
    assert len(data) > 10_000  # a real .pptx with charts/tables is well over 10 KB


def test_report_opens_as_valid_pptx(built_report):
    built_report.seek(0)
    prs = Presentation(built_report)
    assert len(prs.slides) == 8


def test_report_contains_agency_name_and_disclaimer(built_report):
    built_report.seek(0)
    prs = Presentation(built_report)
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    joined = "\n".join(all_text)
    assert "LA Metro" in joined
    assert "independent portfolio project" in joined.lower()
    assert "FTA National Transit Database" in joined


def test_report_labels_real_and_hypothetical_slides(built_report):
    built_report.seek(0)
    prs = Presentation(built_report)
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    joined = "\n".join(all_text)
    assert "REAL DATA" in joined
    assert "HYPOTHETICAL" in joined


def test_report_contains_charts_and_tables(built_report):
    built_report.seek(0)
    prs = Presentation(built_report)
    has_chart = any(shape.has_chart for slide in prs.slides for shape in slide.shapes if shape.shape_type is not None and hasattr(shape, "has_chart"))
    has_table = any(shape.has_table for slide in prs.slides for shape in slide.shapes if hasattr(shape, "has_table"))
    assert has_chart
    assert has_table


def test_report_works_without_any_ai_import():
    # Guard against accidental coupling: report_generator must not import
    # any external AI/LLM SDK.
    import report_generator
    src = Path(report_generator.__file__).read_text()
    for forbidden in ("anthropic", "openai", "ai_advisor"):
        assert forbidden not in src.lower()


def test_report_roadmap_table_has_owner_and_kpi_columns(built_report):
    built_report.seek(0)
    prs = Presentation(built_report)
    roadmap_slide = prs.slides[6]  # Slide 7 (0-indexed): 90-day action plan & roadmap
    tables = [shape.table for shape in roadmap_slide.shapes if shape.has_table]
    assert tables, "Expected the 90-day roadmap table on the action-plan slide"
    header_cells = [cell.text for cell in tables[0].rows[0].cells]
    assert "Workstream" in header_cells
    assert "Owner (role)" in header_cells
    assert "Target KPI" in header_cells
